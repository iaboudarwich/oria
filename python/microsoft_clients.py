"""
Microsoft Graph read helpers for the Oria sidecar.

Covers Outlook mail (read-only Mail.Read), OneDrive files (Files.Read), and
Outlook Calendar (Calendars.Read). Office formats are parsed locally
(python-docx, openpyxl, python-pptx); other binaries go through the existing
extractor router. Content is fetched on demand and returned to the Node app; it
is never persisted here. Access tokens arrive in the signed request body and
are NEVER logged.
"""

from __future__ import annotations

import io
import logging
from typing import Optional

import httpx
from bs4 import BeautifulSoup

logger = logging.getLogger(__name__)

GRAPH = "https://graph.microsoft.com/v1.0"


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


# ── Outlook mail ──────────────────────────────────────────────────────────────

_MAIL_SELECT = "id,subject,from,receivedDateTime,bodyPreview,body,hasAttachments"


def _strip_body(body: dict) -> str:
    content = body.get("content", "") or ""
    if body.get("contentType", "").lower() == "html":
        return BeautifulSoup(content, "html.parser").get_text(" ", strip=True)[:4000]
    return content[:4000]


def outlook_fetch_messages(
    client: httpx.Client, token: str, since_iso: Optional[str], max_messages: int
) -> list[dict]:
    """Fetch recent Outlook messages, newest first, paginating up to max_messages."""
    params = {
        "$select": _MAIL_SELECT,
        "$top": str(min(50, max_messages)),
        "$orderby": "receivedDateTime desc",
    }
    if since_iso:
        params["$filter"] = f"receivedDateTime ge {since_iso}"

    url = f"{GRAPH}/me/messages"
    out: list[dict] = []
    next_link: Optional[str] = None
    while len(out) < max_messages:
        resp = client.get(next_link or url, headers=_auth(token), params=None if next_link else params)
        if resp.status_code == 401:
            raise RuntimeError("Outlook token rejected")
        if resp.status_code != 200:
            break
        data = resp.json()
        for m in data.get("value", []) or []:
            sender = (((m.get("from") or {}).get("emailAddress") or {}).get("address")) or ""
            out.append(
                {
                    "id": m.get("id", ""),
                    "subject": m.get("subject", "") or "",
                    "sender": sender,
                    "date": m.get("receivedDateTime"),
                    "snippet": m.get("bodyPreview", "") or "",
                    "body": _strip_body(m.get("body", {}) or {}),
                    "has_attachments": bool(m.get("hasAttachments")),
                }
            )
            if len(out) >= max_messages:
                break
        next_link = data.get("@odata.nextLink")
        if not next_link:
            break
    return out


# ── OneDrive files ────────────────────────────────────────────────────────────

from extractors.router import route  # noqa: E402  (kept local to this section)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

_ITEM_SELECT = "id,name,file,size,webUrl,lastModifiedDateTime"


def onedrive_get_metadata(client: httpx.Client, token: str, item_id: str) -> Optional[dict]:
    """Fetch one item's metadata. Returns None if inaccessible (403/404)."""
    resp = client.get(
        f"{GRAPH}/me/drive/items/{item_id}",
        headers=_auth(token),
        params={"$select": _ITEM_SELECT},
    )
    if resp.status_code in (403, 404):
        return None
    if resp.status_code != 200:
        return None
    return resp.json()


def _parse_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                out.append(",".join(cells))
    return "\n".join(out)


def _parse_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        out.append(f"# Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text:
                        out.append(text)
    return "\n".join(out)


def onedrive_fetch_text(
    client: httpx.Client, token: str, item_id: str, mime_type: str, name: str
) -> tuple[str, bool]:
    """Download + parse one OneDrive item. Returns (text, accessible). Office
    formats are parsed locally; everything else goes through the extractor."""
    resp = client.get(f"{GRAPH}/me/drive/items/{item_id}/content", headers=_auth(token))
    if resp.status_code in (403, 404):
        return "", False
    if resp.status_code not in (200, 302):
        return "", True
    data = resp.content
    try:
        if mime_type == DOCX_MIME or name.lower().endswith(".docx"):
            return _parse_docx(data), True
        if mime_type == XLSX_MIME or name.lower().endswith(".xlsx"):
            return _parse_xlsx(data), True
        if mime_type == PPTX_MIME or name.lower().endswith(".pptx"):
            return _parse_pptx(data), True
        if mime_type.startswith("text/") or name.lower().endswith((".txt", ".md", ".csv")):
            return data.decode("utf-8", errors="replace"), True
        # PDFs, images, and anything else: the shared extractor (incl. OCR).
        text, _method = route(data, mime_type or "application/octet-stream", name or "file")
        return text, True
    except Exception:
        return "", True


def outlook_calendar_list_events(
    client: httpx.Client, token: str, past_days: int, future_days: int
) -> list[dict]:
    """List Outlook calendar events in [-past_days, +future_days] via calendarView
    (recurring events are expanded into instances). Times are requested in UTC."""
    from datetime import datetime, timedelta, timezone

    now = datetime.now(timezone.utc)
    start = (now - timedelta(days=max(0, past_days))).strftime("%Y-%m-%dT%H:%M:%SZ")
    end = (now + timedelta(days=max(0, future_days))).strftime("%Y-%m-%dT%H:%M:%SZ")
    headers = {**_auth(token), "Prefer": 'outlook.timezone="UTC"'}

    out: list[dict] = []
    url = (
        f"{GRAPH}/me/calendarView?startDateTime={start}&endDateTime={end}"
        "&$select=id,subject,bodyPreview,location,start,end,isAllDay,organizer,attendees,webLink"
        "&$orderby=start/dateTime&$top=250"
    )
    next_link: Optional[str] = None
    while len(out) < 1000:
        resp = client.get(next_link or url, headers=headers)
        if resp.status_code != 200:
            break
        data = resp.json()
        out.extend(data.get("value", []) or [])
        next_link = data.get("@odata.nextLink")
        if not next_link:
            break
    return out


def _iso_utc(dt: Optional[dict]) -> Optional[str]:
    if not dt:
        return None
    val = dt.get("dateTime")
    if not val:
        return None
    return val if val.endswith("Z") else val + "Z"


def outlook_event_dict(item: dict) -> Optional[dict]:
    """Normalize a Graph calendar event into the shared event shape."""
    starts_at = _iso_utc(item.get("start"))
    ends_at = _iso_utc(item.get("end")) or starts_at
    if not starts_at:
        return None
    attendees = [
        {
            "email": (a.get("emailAddress") or {}).get("address", ""),
            "name": (a.get("emailAddress") or {}).get("name"),
        }
        for a in (item.get("attendees") or [])
        if (a.get("emailAddress") or {}).get("address")
    ]
    return {
        "provider_event_id": item.get("id", ""),
        "title": item.get("subject") or "(no title)",
        "description": item.get("bodyPreview"),
        "location": (item.get("location") or {}).get("displayName"),
        "starts_at": starts_at,
        "ends_at": ends_at,
        "is_all_day": bool(item.get("isAllDay")),
        "organizer_email": ((item.get("organizer") or {}).get("emailAddress") or {}).get("address"),
        "attendees": attendees,
        "web_view_link": item.get("webLink"),
    }


def onedrive_meta_dict(item: dict, excerpt: str = "", accessible: bool = True) -> dict:
    """Normalize Graph item metadata into the provider-agnostic file meta shape."""
    size = item.get("size")
    return {
        "provider_file_id": item.get("id", ""),
        "name": item.get("name", "Untitled"),
        "mime_type": (item.get("file") or {}).get("mimeType", "application/octet-stream"),
        "web_view_link": item.get("webUrl"),
        "icon_link": None,
        "thumbnail_link": None,
        "size_bytes": int(size) if isinstance(size, int) else None,
        "modified_time": item.get("lastModifiedDateTime"),
        "excerpt": excerpt,
        "accessible": accessible,
    }
